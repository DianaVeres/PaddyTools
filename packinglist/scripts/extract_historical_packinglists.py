from __future__ import annotations

import hashlib
import json
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path

import openpyxl


def clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def folded(value: object) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", clean(value).upper()) if unicodedata.category(c) != "Mn")


def iso_date(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = clean(value)
    for pattern in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%m-%y"):
        try:
            return datetime.strptime(raw, pattern).date().isoformat()
        except ValueError:
            pass
    return ""


def school_from_sheet(sheet, path: Path) -> str:
    for row in range(1, min(sheet.max_row, 10) + 1):
        for col in range(1, min(sheet.max_column, 8) + 1):
            if folded(sheet.cell(row, col).value) == "COLEGIO":
                for candidate in range(col + 1, min(sheet.max_column, col + 4) + 1):
                    value = clean(sheet.cell(row, candidate).value)
                    if value:
                        return value.upper()
    folder = re.sub(r"^ENVIOS?\s+(A\s+)?", "", folded(path.parent.name)).strip()
    return folder or folded(sheet.title)


def sheet_date(sheet, path: Path) -> str:
    for row in range(1, min(sheet.max_row, 10) + 1):
        for col in range(1, min(sheet.max_column, 8) + 1):
            if folded(sheet.cell(row, col).value) == "FECHA":
                for candidate in range(col + 1, min(sheet.max_column, col + 4) + 1):
                    parsed = iso_date(sheet.cell(row, candidate).value)
                    if parsed:
                        return parsed
    match = re.search(r"(\d{2})-(\d{2})(?:-(\d{4}))?", path.stem)
    if match:
        return f"{match.group(3) or '2026'}-{match.group(2)}-{match.group(1)}"
    return ""


root = Path(sys.argv[1])
output = Path(sys.argv[2])
lists: list[dict] = []
errors: list[str] = []

for path in sorted(root.rglob("*.xlsx")):
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            header_row = 0
            columns: dict[str, int] = {}
            for row in range(1, min(sheet.max_row, 25) + 1):
                values = [folded(sheet.cell(row, col).value) for col in range(1, min(sheet.max_column, 20) + 1)]
                if any("PEDIDO" in value and ("N" in value or "NUM" in value) for value in values):
                    header_row = row
                    for index, value in enumerate(values, 1):
                        if "PEDIDO" in value and ("N" in value or "NUM" in value): columns["order"] = index
                        elif "NOMBRE" in value: columns["name"] = index
                        elif "TIPO" in value: columns["type"] = index
                        elif "FECHA" in value and "ENVIO" in value: columns["shipping"] = index
                        elif "FECHA" in value and "PEDIDO" in value: columns["order_date"] = index
                        elif "TELEFONO" in value: columns["phone"] = index
                        elif "CAJA" in value: columns["box"] = index
                    break
            if not header_row or "order" not in columns:
                continue
            school = school_from_sheet(sheet, path)
            fallback_date = sheet_date(sheet, path)
            orders = []
            current_box = ""
            empty = 0
            for row in range(header_row + 1, sheet.max_row + 1):
                order = re.sub(r"\s+", "", clean(sheet.cell(row, columns["order"]).value).upper())
                if not re.match(r"^(?:PA|IN|R)?\d{3,}$", order):
                    empty += 1
                    if empty >= 30: break
                    continue
                empty = 0
                box = clean(sheet.cell(row, columns.get("box", 8)).value)
                if box: current_box = box
                shipping = iso_date(sheet.cell(row, columns.get("shipping", 5)).value) or fallback_date
                orders.append({
                    "orderNumber": order,
                    "customerName": clean(sheet.cell(row, columns.get("name", 2)).value) or "NO ENCONTRADO",
                    "orderType": clean(sheet.cell(row, columns.get("type", 3)).value) or "ON-LINE",
                    "orderDate": iso_date(sheet.cell(row, columns.get("order_date", 4)).value),
                    "shippingDate": shipping,
                    "phone": clean(sheet.cell(row, columns.get("phone", 7)).value),
                    "box": current_box,
                    "schoolMismatch": False,
                })
            if orders:
                shipping_date = next((o["shippingDate"] for o in orders if o["shippingDate"]), fallback_date)
                signature = school + "|" + shipping_date + "|" + "|".join(sorted(o["orderNumber"] for o in orders))
                lists.append({
                    "id": "HIST-" + hashlib.sha256(signature.encode("utf-8")).hexdigest()[:24],
                    "school": school,
                    "shippingDate": shipping_date,
                    "orderYear": int(shipping_date[:4]),
                    "orderMonth": int(shipping_date[5:7]),
                    "orders": orders,
                    "sourceFile": path.name,
                })
    except Exception as exc:
        errors.append(f"{path.name}: {exc}")

unique = {item["id"]: item for item in lists}
result = sorted(unique.values(), key=lambda item: (item["school"], item["shippingDate"]))
output.parent.mkdir(parents=True, exist_ok=True)
output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({"files": len(list(root.rglob('*.xlsx'))), "packingLists": len(result), "orders": sum(len(x["orders"]) for x in result), "errors": errors}, ensure_ascii=False))
