from __future__ import annotations

import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

root = Path(sys.argv[1])
output = Path(sys.argv[2])
records: list[dict[str, str]] = []

def text(value: object) -> str:
    return '' if value is None else str(value).strip()

def iso_date(value: object) -> str:
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    raw = text(value)
    for pattern in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d %H:%M:%S'):
        try:
            return datetime.strptime(raw, pattern).date().isoformat()
        except ValueError:
            pass
    match = re.search(r'(\d{2})-(\d{2})-(\d{4})', raw)
    return f'{match.group(3)}-{match.group(2)}-{match.group(1)}' if match else raw

for path in sorted(root.rglob('*.xlsx')):
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            school = text(sheet['C3'].value) or path.parent.name.replace('ENVÍOS ', '').replace('ENVÍO ', '')
            header_row = None
            headers: dict[str, int] = {}
            for row_number in range(1, min(sheet.max_row, 20) + 1):
                row = [text(sheet.cell(row_number, column).value).upper() for column in range(1, min(sheet.max_column, 20) + 1)]
                if any('PEDIDO' in value and ('Nº' in value or 'N°' in value or 'NO ' in value) for value in row):
                    header_row = row_number
                    headers = {value: index + 1 for index, value in enumerate(row) if value}
                    break
            if not header_row:
                continue
            order_col = next((column for name, column in headers.items() if 'PEDIDO' in name and ('Nº' in name or 'N°' in name or name.startswith('NO'))), 1)
            shipping_col = next((column for name, column in headers.items() if 'FECHA' in name and ('ENVÍO' in name or 'ENVIO' in name)), 5)
            box_col = next((column for name, column in headers.items() if 'CAJA' in name), 8)
            empty_rows = 0
            max_needed_col = max(order_col, shipping_col, box_col)
            for values in sheet.iter_rows(min_row=header_row + 1, max_col=max_needed_col, values_only=True):
                order = text(values[order_col - 1]).upper().replace(' ', '')
                if not re.match(r'^[A-Z]{0,4}\d{3,}$', order):
                    empty_rows += 1
                    if empty_rows >= 30:
                        break
                    continue
                empty_rows = 0
                shipping_date = iso_date(values[shipping_col - 1]) or iso_date(sheet['C4'].value)
                records.append({'order': order, 'date': shipping_date, 'school': school, 'box': text(values[box_col - 1]), 'file': path.name})
    except Exception:
        continue

deduplicated = list({(item['order'], item['date'], item['school'], item['file']): item for item in records}.values())
output.parent.mkdir(parents=True, exist_ok=True)
output.write_text(json.dumps(deduplicated, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
print(json.dumps({'files': len(list(root.rglob('*.xlsx'))), 'records': len(deduplicated)}, ensure_ascii=True))
